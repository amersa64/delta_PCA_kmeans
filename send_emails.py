"""
Send personalized cold-outreach emails via Gmail SMTP.

Reads buyer_outreach_kit.xlsx (one row per buyer). Sends the email to every
row that has an Email value. Rows with no Email (link-only buyers) are skipped
so you can handle them manually.

Defaults to DRY-RUN so you can preview before sending. Flip DRY_RUN = False
when you're ready to send for real.

Resumable: every attempt is appended to send_log.csv. Re-running the script
skips any email already marked "sent" in the log.
"""

import csv
import smtplib
import time
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook


# === CONFIG — edit these before sending live ===

GMAIL_ADDRESS      = "you@gmail.com"        # your full Gmail address
GMAIL_APP_PASSWORD = ""                      # 16-char App Password (no spaces)
FROM_NAME          = "Your Name"             # what recipients see in the From: line

INPUT_XLSX     = "buyer_outreach_kit.xlsx"
LOG_CSV        = "send_log.csv"
DELAY_SECONDS  = 25                          # pause between sends (Gmail rate-limit friendly)
DRY_RUN        = True                        # flip to False to send for real

# ===============================================


SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 465
REQUIRED_COLUMNS = ("name", "email", "subject", "body")
PLACEHOLDERS = ("[Your name]", "[Your contact link]")


def load_rows(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    missing = [c for c in REQUIRED_COLUMNS if c not in idx]
    if missing:
        raise SystemExit(
            f"{path.name} is missing column(s): {missing}. Found headers: {headers}"
        )

    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if raw is None or not any(raw):
            continue
        rows.append({
            c: (str(raw[idx[c]]).strip() if raw[idx[c]] is not None else "")
            for c in REQUIRED_COLUMNS
        })
    return rows


def load_sent(log_path: Path):
    if not log_path.exists():
        return set()
    sent = set()
    with log_path.open(newline="") as f:
        for row in csv.DictReader(f):
            if row.get("status") == "sent":
                sent.add(row["email"].lower())
    return sent


def append_log(log_path: Path, email: str, status: str, note: str = ""):
    new = not log_path.exists()
    with log_path.open("a", newline="") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["timestamp", "email", "status", "note"])
        w.writerow([time.strftime("%Y-%m-%d %H:%M:%S"), email, status, note])


def check_placeholders(rows):
    bad = []
    for r in rows:
        blob = r["subject"] + " " + r["body"]
        if any(p in blob for p in PLACEHOLDERS):
            bad.append(r["email"] or r["name"])
    return bad


def build_message(row):
    msg = EmailMessage()
    msg["From"]    = f"{FROM_NAME} <{GMAIL_ADDRESS}>"
    msg["To"]      = row["email"]
    msg["Subject"] = row["subject"]
    msg.set_content(row["body"])
    return msg


def main():
    here = Path(__file__).resolve().parent
    xlsx = here / INPUT_XLSX
    log_path = here / LOG_CSV

    if not xlsx.exists():
        raise SystemExit(f"Can't find {xlsx}. Run make_templates.py first or drop your file there.")
    if "@" not in GMAIL_ADDRESS or GMAIL_ADDRESS.startswith("you@"):
        raise SystemExit("Edit GMAIL_ADDRESS at the top of this file.")
    if not DRY_RUN and not GMAIL_APP_PASSWORD:
        raise SystemExit("Set GMAIL_APP_PASSWORD (16-char Gmail App Password) before sending live.")

    rows = load_rows(xlsx)
    already_sent = load_sent(log_path)

    with_email = [r for r in rows if r["email"]]
    no_email   = [r for r in rows if not r["email"]]
    todo       = [r for r in with_email if r["email"].lower() not in already_sent]

    bad = check_placeholders(todo)

    print(f"Input file        : {xlsx.name}")
    print(f"Total rows        : {len(rows)}")
    print(f"With email        : {len(with_email)}  (auto-sendable)")
    print(f"Link-only (manual): {len(no_email)}")
    print(f"Already sent      : {len(with_email) - len(todo)}")
    print(f"To send this run  : {len(todo)}")
    print(f"Mode              : {'DRY-RUN (no emails sent)' if DRY_RUN else 'LIVE'}")
    print()

    if bad:
        print(f"WARNING: {len(bad)} row(s) still contain placeholder text "
              f"({', '.join(PLACEHOLDERS)}):")
        for b in bad[:10]:
            print(f"  - {b}")
        if len(bad) > 10:
            print(f"  ... and {len(bad) - 10} more")
        print("Do a find-and-replace in the XLSX, save, and re-run.")
        if not DRY_RUN:
            raise SystemExit("Refusing to send live until placeholders are filled in.")
        print()

    if DRY_RUN:
        preview = todo[:3]
        for r in preview:
            print("-" * 60)
            print(f"To:      {r['name']} <{r['email']}>")
            print(f"Subject: {r['subject']}")
            print()
            body = r["body"]
            print(body if len(body) <= 600 else body[:600] + "\n... [truncated]")
        print("-" * 60)
        print(f"Showed {len(preview)} of {len(todo)} preview message(s).")
        print("Dry-run complete. Set DRY_RUN = False at the top of this file to send.")
        return

    print(f"Sending live. {DELAY_SECONDS}s between emails. Ctrl-C to abort.")
    print("Starting in 5 seconds...")
    time.sleep(5)

    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT) as server:
        server.login(GMAIL_ADDRESS, GMAIL_APP_PASSWORD)
        for i, r in enumerate(todo, start=1):
            label = f"[{i}/{len(todo)}] {r['email']}"
            try:
                server.send_message(build_message(r))
                append_log(log_path, r["email"], "sent")
                print(f"{label}  sent")
            except Exception as e:
                append_log(log_path, r["email"], "error", repr(e))
                print(f"{label}  ERROR: {e}")
            if i < len(todo):
                time.sleep(DELAY_SECONDS)

    print(f"\nDone. Results logged to {log_path.name}.")
    if no_email:
        print(f"Reminder: {len(no_email)} link-only buyer(s) still need manual outreach.")


if __name__ == "__main__":
    main()
